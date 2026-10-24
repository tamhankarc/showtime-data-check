"""Transient normalization, deduplication, filtering and XLSX generation."""
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
import json
import re
from urllib.parse import urlsplit
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import select
from .models import FormatMapping, MovieFormatOverride, MovieGroup, ProviderMovie, ProviderTheater, Theater


def text_key(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def language_value(value):
    if isinstance(value, (list, tuple, set)):
        return ", ".join(sorted({str(item).strip().upper() for item in value if item}))
    return str(value or "").strip().upper()


def usable_url(value):
    if not isinstance(value, str) or len(value) > 8192 or any(ord(c) < 32 for c in value):
        return None
    try:
        parsed = urlsplit(value.strip())
        return value.strip() if parsed.scheme in {"http", "https"} and parsed.hostname and not parsed.username and not parsed.password else None
    except ValueError:
        return None


def normalize_rows(db, connection, source_rows):
    movie_rows = {row.external_id: row for row in db.scalars(select(ProviderMovie).where(ProviderMovie.connection_id == connection.id))}
    groups = {row.id: row for row in db.scalars(select(MovieGroup).where(MovieGroup.country_code == connection.country_code))}
    observed = {row.external_id: row for row in db.scalars(select(ProviderTheater).where(ProviderTheater.connection_id == connection.id))}
    canonical = {row.id: row for row in db.scalars(select(Theater).where(Theater.country_code == connection.country_code))}
    mappings = {(row.country_code, text_key(row.raw_format)): row for row in db.scalars(select(FormatMapping).where(FormatMapping.provider_code == connection.provider_code))}
    overrides = {(row.group_id, row.format_mapping_id): row.published_format for row in db.scalars(select(MovieFormatOverride).where(MovieFormatOverride.active.is_(True)))}
    output = []
    for source in source_rows:
        movie = movie_rows.get(str(source.get("title")))
        provider_theater = observed.get(str(source.get("theater")))
        if provider_theater and provider_theater.review_status == "ignored":
            continue
        theater = canonical.get(provider_theater.theater_id) if provider_theater and provider_theater.review_status == "confirmed" else None
        if movie is None or (theater and not theater.active):
            continue
        local_date, local_time = str(source.get("date") or "")[:10], str(source.get("time") or "")[:8]
        zone = theater.timezone if theater else connection.timezone
        try:
            if date.fromisoformat(local_date) < datetime.now(ZoneInfo(zone)).date():
                continue
            datetime.strptime(local_time[:5], "%H:%M")
        except (ValueError, TypeError):
            continue
        raw = str(source.get("format") or "2D")[:1000]
        mapping = mappings.get((connection.country_code, text_key(raw))) or mappings.get(("", text_key(raw)))
        published = overrides.get((movie.group_id, mapping.id), mapping.published_format) if mapping else raw
        audio = language_value(source.get("audio") or (mapping.audio_language if mapping else ""))
        subtitle = language_value(source.get("subtitle") or (mapping.subtitle_language if mapping else ""))
        semantic = text_key(raw + " " + published)
        vip = bool(source.get("vip") or (mapping and mapping.is_luxury) or re.search(r"\b(vip|luxo|luxury|prime|premier)\b", semantic))
        dubbed = bool((mapping and mapping.is_dubbed) or re.search(r"\b(dubbed|dublado|dub)\b", semantic))
        subtitled = bool((mapping and mapping.is_subtitled) or subtitle or re.search(r"\b(subtitled|legendado|leg)\b", semantic))
        links = []
        for link in source.get("links") or []:
            url = usable_url(link.get("url"))
            if url:
                links.append({"name": str(link.get("name") or connection.provider_code), "url": url, "kind": link.get("kind") or "booking"})
        if not links:
            for key, kind in (("booking_url", "booking"), ("movie_url", "movie")):
                if usable_url(source.get(key)):
                    links.append({"name": connection.provider_code if kind == "booking" else "Movie page", "url": usable_url(source[key]), "kind": kind})
                    break
        if not links and theater and usable_url(theater.website):
            links.append({"name": "Theater website", "url": usable_url(theater.website), "kind": "theater"})
        tid = theater.id if theater else f"{connection.provider_code}:{source.get('theater')}"
        group = groups.get(movie.group_id)
        output.append({"id": f"{connection.provider_code}:{source.get('id', '')}", "movie_name": group.name if group else movie.name,
                       "movie_group_id": movie.group_id, "provider_movie_id": movie.external_id, "provider_movie_name": movie.name,
                       "theater_id": tid, "theater_name": theater.name if theater else (provider_theater.name if provider_theater else str(source.get("theater"))),
                       "chain": theater.chain if theater else (provider_theater.chain if provider_theater else None),
                       "address": theater.address if theater else (provider_theater.address if provider_theater else None),
                       "city": theater.city if theater else (provider_theater.city if provider_theater else None),
                       "country_code": connection.country_code, "date": local_date, "time": local_time[:5], "timezone": zone,
                       "published_format": published, "raw_format": raw, "format_mapped": bool(mapping), "audio_language": audio,
                       "subtitle_language": subtitle, "vip": vip, "dubbed": dubbed, "subtitled": subtitled,
                       "providers": [connection.provider_code], "links": links, "has_usable_link": bool(links),
                       "imax": bool(re.search(r"\bIMAX\b", published, re.I)), "provider_ids": [str(source.get("id", ""))]})
    return output


def consolidate(rows):
    merged = {}
    for row in rows:
        # A manually selected BR ID set is one logical check. IDs across groups may refer to localized names.
        key = (str(row["theater_id"]), row["date"], row["time"], text_key(row["published_format"]),
               text_key(row["audio_language"]), text_key(row["subtitle_language"]), row["vip"], row["dubbed"], row["subtitled"])
        if key not in merged:
            merged[key] = {**row, "links": list(row["links"]), "providers": list(row["providers"]),
                           "raw_formats": [row["raw_format"]], "provider_ids": list(row["provider_ids"]),
                           "movie_names": [row["movie_name"]], "movie_group_ids": [row["movie_group_id"]]}
            continue
        existing = merged[key]
        for field, values in (("links", row["links"]), ("providers", row["providers"]), ("raw_formats", [row["raw_format"]]),
                              ("provider_ids", row["provider_ids"]), ("movie_names", [row["movie_name"]]), ("movie_group_ids", [row["movie_group_id"]])):
            for value in values:
                if value not in existing[field]:
                    existing[field].append(value)
        existing["has_usable_link"] = bool(existing["links"])
        existing["raw_format"] = " / ".join(existing["raw_formats"])
        existing["movie_name"] = " / ".join(existing["movie_names"])
    return sorted(merged.values(), key=lambda r: (r["date"], r["time"], text_key(r["theater_name"])))


def filtered(rows, filters):
    output = []
    query = text_key(filters.get("q"))
    exact = str(filters.get("exact", "false")).lower() == "true"
    for row in rows:
        if filters.get("chain") not in (None, "", "all", row.get("chain")):
            continue
        if filters.get("format") not in (None, "", "all", row["published_format"]):
            continue
        if filters.get("language") not in (None, "", "all", row["audio_language"], row["subtitle_language"]):
            continue
        if filters.get("date_from") and row["date"] < filters["date_from"]:
            continue
        if filters.get("links") in ("missing", "no") and row["has_usable_link"]:
            continue
        if filters.get("links") in ("available", "yes") and not row["has_usable_link"]:
            continue
        values = [row.get(k) for k in ("movie_name", "theater_name", "theater_id", "address", "city", "chain", "raw_format", "provider_movie_id")]
        values += row.get("provider_ids", [])
        if query and not any(query == text_key(v) if exact else query in text_key(v) for v in values):
            continue
        output.append(row)
    return output


def theater_rows(rows):
    groups = defaultdict(list)
    for row in rows:
        groups[str(row["theater_id"])].append(row)
    output = []
    for values in groups.values():
        first, last = min(values, key=lambda r: (r["date"], r["time"])), max(values, key=lambda r: (r["date"], r["time"]))
        output.append({**{key: first[key] for key in ("theater_id", "theater_name", "chain", "address", "city", "country_code")},
                       "id": str(first["theater_id"]), "movie_name": " / ".join(sorted({r["movie_name"] for r in values})),
                       "showtimes": len(values), "first": f"{first['date']} {first['time']}", "last": f"{last['date']} {last['time']}",
                       "formats": sorted({r["published_format"] for r in values}), "providers": sorted({p for r in values for p in r["providers"]})})
    return sorted(output, key=lambda row: text_key(row["theater_name"]))


def metrics(rows):
    return {"showtimes": len(rows), "theaters": len({str(row["theater_id"]) for row in rows}),
            "no_usable_link": sum(not row["has_usable_link"] for row in rows), "imax": sum(row["imax"] for row in rows)}


def xlsx_export(rows, metadata, view):
    book = Workbook()
    sheet = book.active
    sheet.title = "Theaters" if view == "theaters" else "Showtimes"
    fields = (["movie_name", "theater_id", "theater_name", "chain", "address", "city", "country_code", "showtimes", "first", "last", "formats", "providers"]
              if view == "theaters" else ["movie_name", "theater_id", "theater_name", "chain", "address", "city", "country_code", "date", "time", "timezone", "published_format", "raw_format", "audio_language", "subtitle_language", "providers", "links", "has_usable_link"])
    def safe(value):
        if isinstance(value, (list, dict)):
            value = json.dumps(value, ensure_ascii=False)
        if isinstance(value, str):
            value = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", value)
            if value[:1] in "=+-@":
                return "'" + value
        return value
    sheet.append([field.replace("_", " ").title() for field in fields])
    for row in rows:
        sheet.append([safe(row.get(field)) for field in fields])
    sheet.freeze_panes, sheet.auto_filter.ref = "A2", sheet.dimensions
    for cell in sheet[1]:
        cell.fill, cell.font = PatternFill("solid", fgColor="173B63"), Font(color="FFFFFF", bold=True)
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(55, max(14, len(str(column[0].value)) + 4))
    meta = book.create_sheet("Check details")
    meta.append(["Field", "Value"])
    for key, value in metadata.items():
        meta.append([key, safe(value)])
    meta.column_dimensions["A"].width, meta.column_dimensions["B"].width = 26, 100
    for row in meta:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    result = BytesIO()
    book.save(result)
    return result.getvalue()
