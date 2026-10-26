"""Offline end-to-end API/worker tests; only temporary SQLite fixtures are written."""
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
import json
import threading
import time
from types import SimpleNamespace

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from sqlalchemy import event, inspect, select

from app.config import Settings
from app.db import make_engine
from app.main import create_app
from app.models import (Base, Connection, Country, FormatMapping, Job, MovieGroup,
                        ProviderMovie, ProviderTheater, Theater)
from app.providers import ProviderError, ProviderResult
from app.security import password_hash


PASSWORD = "fixture-password-only"
DATE = "2099-11-02"
TERMINAL = {"done", "partial", "failed", "cancelled", "expired", "stale"}


@pytest.fixture(scope="module")
def hashes():
    return password_hash(PASSWORD, salt="offline-administrator"), password_hash(PASSWORD, salt="offline-viewer")


class FixtureProvider:
    def __init__(self):
        self.calls = []
        self.lock = threading.Lock()
        self.fail_catalog = set()
        self.fail_showtimes = set()
        self.movie_overrides = {}
        self.theater_overrides = {}
        self.block_codes = set()
        self.started = threading.Event()
        self.release = threading.Event()

    def _begin(self, kind, connection, external_ids, cancel):
        code = connection["provider_code"]
        with self.lock:
            self.calls.append((kind, code, sorted(external_ids or [])))
        if code in self.block_codes:
            self.started.set()
            while not self.release.wait(0.01):
                if cancel.is_set():
                    raise ProviderError("Fixture operation cancelled.")
        failures = self.fail_catalog if kind == "catalog" else self.fail_showtimes
        if code in failures:
            raise ProviderError("The fixture provider is temporarily unavailable.")
        return code

    @staticmethod
    def _movies(code):
        if code == "INGRESSO":
            return [{"id": "ing-m", "name": "Feature from provider"}, {"id": "ing-alt", "name": "Another movie"}]
        if code == "CINEMARK":
            return [{"id": "cin-m", "name": "Feature localized"}]
        return [{"id": "is-m", "name": "IS Movie"}]

    @staticmethod
    def _theaters(code):
        if code == "INGRESSO":
            return [{"id": "ing-t", "name": "Provider cinema name", "address": "Observed provider address",
                     "city": "Rio", "timezone": "America/Sao_Paulo", "chain_name": "Observed Chain"},
                    {"id": "ing-empty", "name": "Unlinked cinema", "city": "Rio", "timezone": "America/Sao_Paulo"}]
        return [{"id": "cin-t", "name": "Cinemark provider name", "city": "Rio", "timezone": "America/Sao_Paulo"}]

    def catalog(self, connection, cancel):
        code = self._begin("catalog", connection, None, cancel)
        return ProviderResult(movies=deepcopy(self.movie_overrides.get(code, self._movies(code))),
                              theaters=deepcopy(self.theater_overrides.get(code, self._theaters(code))))

    def showtimes(self, connection, external_ids, cancel):
        code = self._begin("check", connection, external_ids, cancel)
        title, theater = ("ing-m", "ing-t") if code == "INGRESSO" else ("cin-m", "cin-t")
        common = {"title": title, "theater": theater, "date": DATE, "time": "14:00", "format": "IMAX | Dublado",
                  "audio": ["PT"] if code == "INGRESSO" else [], "subtitle": [], "attributes": [],
                  "booking_url": f"https://booking.example/{code.lower()}", "movie_url": None,
                  "links": [{"name": code, "url": f"https://booking.example/{code.lower()}", "kind": "booking"}]}
        rows = [{**common, "id": f"{code}-normal"}] if title in external_ids else []
        if code == "INGRESSO" and title in external_ids:
            rows += [{**common, "id": "INGRESSO-vip", "format": "IMAX | Dublado | Vip", "vip": True,
                      "links": [{"name": "Ingresso VIP", "url": "https://booking.example/vip", "kind": "booking"}]},
                     {**common, "id": "INGRESSO-missing", "theater": "ing-empty", "time": "15:00",
                      "format": "Unknown Laser", "audio": [], "links": [], "booking_url": None}]
        return ProviderResult(movies=self._movies(code), theaters=self._theaters(code), showtimes=rows)


@pytest.fixture
def env(tmp_path, hashes):
    database_url = "sqlite:///" + (tmp_path / "integration.sqlite").as_posix()
    settings = Settings(database_url=database_url, environment="development", app_origin="http://localhost:3001", proxy_secret="fixture-proxy-secret",
        admin_username="administrator", admin_password_hash=hashes[0],
        viewer_username="viewer", viewer_password_hash=hashes[1], snapshot_seconds=900)
    engine, provider = make_engine(database_url), FixtureProvider()
    app = create_app(settings, provider, engine)
    with TestClient(app) as client:
        client.headers.update({"x-proxy-secret": settings.proxy_secret, "origin": settings.app_origin})
        with app.state.sessions() as db:
            db.add_all([Country(code="BR", name="Brazil", timezone="America/Sao_Paulo"),
                        Country(code="AU", name="Australia", timezone="Australia/Sydney")])
            db.flush()
            db.add_all([Connection(id=1, country_code="BR", provider_code="INGRESSO", provider_country_code="BR", timezone="America/Sao_Paulo"),
                        Connection(id=2, country_code="BR", provider_code="CINEMARK", provider_country_code="BR", timezone="America/Sao_Paulo")])
            db.add_all([MovieGroup(id=201, country_code="BR", name="Feature"),
                        MovieGroup(id=202, country_code="BR", name="Localized Feature"),
                        MovieGroup(id=203, country_code="BR", name="Another movie")])
            db.add_all([Theater(id=101, country_code="BR", name="Canonical Cinema", address="Manual address", city="Rio",
                                website="https://theater.example/", chain="Cinemark", timezone="America/Sao_Paulo"),
                        Theater(id=102, country_code="BR", name="No Link Cinema", address="Other street", city="Rio",
                                chain="Independent", timezone="America/Sao_Paulo")])
            db.flush()
            db.add_all([ProviderMovie(id=301, connection_id=1, external_id="ing-m", name="Feature", group_id=201, manually_grouped=True),
                        ProviderMovie(id=302, connection_id=2, external_id="cin-m", name="Localized Feature", group_id=202),
                        ProviderMovie(id=303, connection_id=1, external_id="ing-alt", name="Another movie", group_id=203)])
            db.add_all([ProviderTheater(id=401, connection_id=1, external_id="ing-t", name="Ingresso Cinema", theater_id=101, review_status="confirmed"),
                        ProviderTheater(id=402, connection_id=2, external_id="cin-t", name="Cinemark Cinema", theater_id=101, review_status="confirmed"),
                        ProviderTheater(id=403, connection_id=1, external_id="ing-empty", name="No Link Cinema", theater_id=102, review_status="confirmed")])
            for identifier, code in ((501, "INGRESSO"), (502, "CINEMARK")):
                db.add(FormatMapping(id=identifier, country_code="BR", provider_code=code, raw_format="IMAX | Dublado",
                    published_format="IMAX", audio_language="PT", is_dubbed=True))
            db.add(FormatMapping(id=503, country_code="BR", provider_code="INGRESSO", raw_format="IMAX | Dublado | Vip",
                published_format="IMAX", audio_language="PT", is_dubbed=True, is_luxury=True))
            db.commit()
        context = SimpleNamespace(client=client, app=app, settings=settings, provider=provider, engine=engine)
        yield context
        provider.release.set()
        app.state.jobs.executor.shutdown(wait=True, cancel_futures=True)
    engine.dispose()


def login(env, role="admin"):
    username = env.settings.admin_username if role == "admin" else env.settings.viewer_username
    response = env.client.post("/api/login", json={"username": username, "password": PASSWORD})
    assert response.status_code == 200, response.text
    assert response.json()["role"] == role
    return response


def wait_job(env, identifier, expected=None):
    deadline = time.monotonic() + 8
    while time.monotonic() < deadline:
        response = env.client.get(f"/api/jobs/{identifier}")
        assert response.status_code == 200, response.text
        job = response.json()
        if job["status"] in TERMINAL:
            if expected is not None:
                assert job["status"] == expected, job
            return job
        time.sleep(0.01)
    raise AssertionError(f"Fixture job {identifier} did not finish.")


def start_check(env, expected="done", provider_ids=None):
    response = env.client.post("/api/checks", json={"country_code": "BR", "provider_movie_ids": provider_ids or [301, 302]})
    assert response.status_code == 200, response.text
    identifier = response.json()["job_id"]
    wait_job(env, identifier, expected)
    return identifier


def refresh(env, identifiers=None, expected="done"):
    response = env.client.post("/api/catalogs/refresh", json={"connection_ids": identifiers or [1, 2]})
    assert response.status_code == 200, response.text
    identifier = response.json()["job_id"]
    wait_job(env, identifier, expected)
    return identifier


def test_proxy_secret_csrf_and_authentication_are_required(env):
    assert env.client.get("/health", headers={"x-proxy-secret": ""}).status_code == 200
    assert env.client.get("/api/bootstrap", headers={"x-proxy-secret": ""}).status_code == 403
    assert env.client.get("/api/bootstrap").status_code == 401
    assert env.client.post("/api/login", headers={"origin": "https://attacker.invalid"},
        json={"username": "administrator", "password": PASSWORD}).status_code == 403
    assert env.client.post("/api/login", json={"username": "administrator", "password": "incorrect"}).status_code == 401
    response = login(env)
    cookie = response.headers.get("set-cookie", "").lower()
    assert "httponly" in cookie and "samesite=strict" in cookie
    assert response.headers["cache-control"] == "no-store, private"
    assert env.client.get("/api/session").json()["role"] == "admin"
    assert env.client.post("/api/logout").status_code == 200
    assert env.client.get("/api/session").status_code == 401
    assert env.provider.calls == []


@pytest.mark.parametrize("method,path,payload", [
    ("post", "/api/countries", {"code": "NZ", "name": "New Zealand", "timezone": "Pacific/Auckland"}),
    ("patch", "/api/theaters/101", {"name": "Unauthorised edit"}),
    ("patch", "/api/provider-theaters/401", {"theater_id": 102, "review_status": "confirmed"}),
    ("post", "/api/movie-groups/move", {"provider_movie_ids": [302], "target_group_id": 201}),
    ("delete", "/api/format-mappings/501", None),
])
def test_viewer_cannot_mutate_reference_data(env, method, path, payload):
    login(env, "viewer")
    response = env.client.request(method, path, json=payload) if payload is not None else env.client.request(method, path)
    assert response.status_code == 403
    assert env.provider.calls == []


def test_reads_and_initial_page_data_never_download_provider_data(env):
    login(env, "viewer")
    for path in ("/api/bootstrap", "/api/movies?country=BR", "/api/provider-movies?country=BR",
                 "/api/theaters?country=BR", "/api/provider-theaters?country=BR", "/api/format-mappings"):
        assert env.client.get(path).status_code == 200, path
    assert env.provider.calls == []


def test_viewer_can_explicitly_refresh_and_load_current_data(env):
    login(env, "viewer")
    refresh(env)
    identifier = start_check(env)
    result = env.client.get(f"/api/checks/{identifier}").json()
    assert result["status"] == "done" and result["total"] > 0
    assert any(call[0] == "catalog" for call in env.provider.calls)
    assert any(call[0] == "check" for call in env.provider.calls)


def test_br_multiple_provider_ids_merge_links_and_keep_vip_distinct(env):
    login(env)
    identifier = start_check(env)
    response = env.client.get(f"/api/checks/{identifier}")
    assert response.status_code == 200
    result = response.json()
    assert result["metrics"] == {"showtimes": 3, "theaters": 2, "no_usable_link": 1, "imax": 2}
    combined = next(row for row in result["rows"] if len(row["providers"]) == 2)
    assert set(combined["providers"]) == {"INGRESSO", "CINEMARK"}
    assert len(combined["links"]) == 2 and combined["vip"] is False
    vip = next(row for row in result["rows"] if row["vip"])
    assert vip["time"] == combined["time"] and vip["theater_id"] == combined["theater_id"]
    assert vip["providers"] == ["INGRESSO"]
    assert ("check", "INGRESSO", ["ing-m"]) in env.provider.calls
    assert ("check", "CINEMARK", ["cin-m"]) in env.provider.calls


def test_mapping_fallback_languages_and_provider_local_times(env):
    login(env)
    identifier = start_check(env)
    rows = env.client.get(f"/api/checks/{identifier}").json()["rows"]
    mapped = next(row for row in rows if len(row["providers"]) == 2)
    assert mapped["published_format"] == "IMAX" and mapped["audio_language"] == "PT"
    assert mapped["dubbed"] is True and mapped["subtitled"] is False
    unknown = next(row for row in rows if row["raw_format"] == "Unknown Laser")
    assert unknown["published_format"] == "Unknown Laser" and unknown["format_mapped"] is False
    assert unknown["audio_language"] == "" and unknown["subtitle_language"] == ""
    assert mapped["date"] == DATE and mapped["time"] == "14:00"
    assert mapped["timezone"] == "America/Sao_Paulo"


def test_theater_view_metrics_and_filter_changes_reuse_snapshot_without_requests(env):
    login(env)
    identifier = start_check(env)
    before = list(env.provider.calls)
    result = env.client.get(f"/api/checks/{identifier}?view=theaters&chain=Cinemark").json()
    assert result["total"] == 1 and result["rows"][0]["showtimes"] == 2
    assert result["rows"][0]["providers"] == ["CINEMARK", "INGRESSO"]
    assert result["metrics"]["showtimes"] == 2
    result = env.client.get(f"/api/checks/{identifier}?q=No+Link+Cinema&exact=true").json()
    assert result["total"] == 1 and result["metrics"]["no_usable_link"] == 1
    assert env.provider.calls == before


def test_export_contains_current_filtered_rows_and_completeness_metadata(env):
    login(env)
    identifier = start_check(env)
    response = env.client.get(f"/api/checks/{identifier}/export?view=showtimes&chain=Cinemark")
    assert response.status_code == 200 and response.content.startswith(b"PK")
    assert "showtimes-BR-" in response.headers["content-disposition"]
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Showtimes", "Check details"]
    sheet = book["Showtimes"]
    assert sheet.max_row == 3
    columns = [cell.value for cell in sheet[1]]
    assert "Audio Language" in columns and "Links" in columns
    for row in sheet.iter_rows(min_row=2, values_only=True):
        assert row[columns.index("Chain")] == "Cinemark"
    details = dict(book["Check details"].iter_rows(min_row=2, values_only=True))
    assert details["status"] == "done"
    assert json.loads(details["selection"])["provider_movie_ids"] == [301, 302]
    assert json.loads(details["filters"])["chain"] == "Cinemark"
    assert all(value["fetched_at"] for value in json.loads(details["provider_outcomes"]))


def test_expiry_disables_result_and_export_without_refreshing(env):
    login(env)
    identifier = start_check(env)
    before = list(env.provider.calls)
    with env.app.state.jobs.lock:
        snapshot = env.app.state.jobs.snapshots[identifier]
        for slice_value in snapshot.slices.values():
            slice_value.deadline = time.monotonic() - 1
    assert env.client.get(f"/api/checks/{identifier}").status_code == 410
    assert env.client.get(f"/api/checks/{identifier}/export").status_code == 410
    assert env.provider.calls == before
    assert env.app.state.jobs.snapshots[identifier].slices == {}


def test_mapping_change_marks_result_stale_and_disables_export(env):
    login(env)
    identifier = start_check(env)
    before = list(env.provider.calls)
    response = env.client.patch("/api/format-mappings/501", json={"published_format": "IMAX Premium"})
    assert response.status_code == 200
    assert env.client.get(f"/api/checks/{identifier}").status_code == 410
    assert env.client.get(f"/api/checks/{identifier}/export").status_code == 410
    assert env.provider.calls == before


def test_partial_provider_failure_retry_does_not_extend_successful_slice(env):
    login(env)
    env.provider.fail_showtimes.add("CINEMARK")
    identifier = start_check(env, "partial")
    result = env.client.get(f"/api/checks/{identifier}").json()
    assert result["status"] == "partial"
    assert {value["status"] for value in result["provider_outcomes"]} == {"done", "failed"}
    with env.app.state.jobs.lock:
        success = env.app.state.jobs.snapshots[identifier].slices["INGRESSO"]
        original = (success.deadline, success.expires_at, success.fetched_at)
    env.provider.fail_showtimes.clear()
    before = len(env.provider.calls)
    response = env.client.post(f"/api/checks/{identifier}/retry", json={"provider_codes": ["CINEMARK"]})
    assert response.status_code == 200
    wait_job(env, identifier, "done")
    assert env.provider.calls[before:] == [("check", "CINEMARK", ["cin-m"])]
    with env.app.state.jobs.lock:
        retained = env.app.state.jobs.snapshots[identifier].slices["INGRESSO"]
        assert (retained.deadline, retained.expires_at, retained.fetched_at) == original
    assert env.client.get(f"/api/checks/{identifier}").json()["metrics"]["showtimes"] == 3


def test_failed_catalog_keeps_previous_movies_and_confirmed_assignments(env):
    login(env)
    with env.app.state.sessions() as db:
        connection = db.get(Connection, 1)
        connection.last_success_at = datetime(2026, 1, 1, tzinfo=timezone.utc)
        db.commit()
    env.provider.fail_catalog.add("INGRESSO")
    refresh(env, [1], "failed")
    with env.app.state.sessions() as db:
        assert db.get(ProviderMovie, 301).active is True
        assert db.get(ProviderMovie, 301).group_id == 201
        assert db.get(ProviderMovie, 301).manually_grouped is True
        assert db.get(ProviderTheater, 401).theater_id == 101
        assert db.get(ProviderTheater, 401).review_status == "confirmed"
        assert db.get(Connection, 1).last_success_at.year == 2026
        assert db.get(Connection, 1).last_status == "failed"


def test_successful_catalog_preserves_manual_theater_edits_and_assignments(env):
    login(env)
    refresh(env, [1])
    with env.app.state.sessions() as db:
        observed, canonical = db.get(ProviderTheater, 401), db.get(Theater, 101)
        assert observed.theater_id == 101 and observed.review_status == "confirmed"
        assert observed.name == "Provider cinema name"
        assert canonical.name == "Canonical Cinema" and canonical.address == "Manual address"
        assert db.get(ProviderMovie, 301).group_id == 201
        assert db.get(ProviderMovie, 301).manually_grouped is True


def test_new_provider_theater_needs_review_instead_of_name_based_merge(env):
    login(env)
    env.provider.theater_overrides["INGRESSO"] = [{"id": "new-t", "name": "Canonical Cinema", "address": "Manual address",
        "city": "Rio", "timezone": "America/Sao_Paulo"}]
    refresh(env, [1])
    with env.app.state.sessions() as db:
        observed = db.scalar(select(ProviderTheater).where(ProviderTheater.connection_id == 1, ProviderTheater.external_id == "new-t"))
        assert observed is not None and observed.review_status == "needs_review"
        assert observed.theater_id != 101


def test_admin_can_group_split_and_reassign_imported_brazil_ids(env):
    login(env)
    response = env.client.post("/api/movie-groups/move", json={"provider_movie_ids": [302], "target_group_id": 201})
    assert response.status_code == 200 and response.json()["group_id"] == 201
    with env.app.state.sessions() as db:
        assert db.get(ProviderMovie, 302).group_id == 201 and db.get(ProviderMovie, 302).manually_grouped is True
        assert db.get(MovieGroup, 202).active is False
    response = env.client.post("/api/movie-groups/move", json={"provider_movie_ids": [302]})
    assert response.status_code == 200 and response.json()["group_id"] != 201
    with env.app.state.sessions() as db:
        assert db.get(ProviderMovie, 302).group_id == response.json()["group_id"]
    assert env.provider.calls == []


def test_admin_theater_assignment_invalidates_existing_results(env):
    login(env)
    identifier = start_check(env)
    response = env.client.patch("/api/provider-theaters/402", json={"theater_id": 102, "review_status": "confirmed"})
    assert response.status_code == 200
    assert env.client.get(f"/api/checks/{identifier}").status_code == 410
    with env.app.state.sessions() as db:
        assert db.get(ProviderTheater, 402).theater_id == 102


def test_cancel_discards_snapshot_and_no_showtime_table_or_job_body_is_durable(env):
    login(env)
    identifier = start_check(env)
    assert env.client.delete(f"/api/checks/{identifier}").status_code == 200
    assert env.client.get(f"/api/checks/{identifier}").status_code == 410
    assert not any("showtime" in name or "session_snapshot" in name for name in inspect(env.engine).get_table_names())
    with env.app.state.sessions() as db:
        job = db.get(Job, identifier)
        durable = json.dumps({"parameters": job.parameters, "outcomes": job.provider_outcomes})
        assert "INGRESSO-normal" not in durable and "https://booking.example/" not in durable
        assert "raw_format" not in durable


def test_invalid_reference_changes_rejected_and_generic_initial_is_local(env):
    login(env)
    assert env.client.patch("/api/theaters/101", json={"latitude": 100}).status_code == 400
    assert env.client.patch("/api/theaters/101", json={"website": "javascript:alert(1)"}).status_code == 400
    assert env.client.patch("/api/countries/BR", json={"code": "XX"}).status_code == 400
    response = env.client.get("/api/theaters?country=BR")
    assert {row["generic_initial"] for row in response.json()["items"]} == {"c", "n"}
    assert env.provider.calls == []


def test_running_check_can_be_cancelled_without_retaining_partial_rows(env):
    login(env)
    env.provider.block_codes = {"INGRESSO", "CINEMARK"}
    response = env.client.post("/api/checks", json={"country_code": "BR", "provider_movie_ids": [301, 302]})
    assert response.status_code == 200
    identifier = response.json()["job_id"]
    assert env.provider.started.wait(3)
    assert env.client.delete(f"/api/checks/{identifier}").status_code == 200
    env.provider.release.set()
    wait_job(env, identifier, "cancelled")
    assert env.client.get(f"/api/checks/{identifier}/export").status_code == 410
    assert env.app.state.jobs.snapshots[identifier].slices == {}


def test_failed_checks_without_bodies_do_not_exhaust_retained_snapshot_capacity(env):
    login(env)
    env.settings.max_snapshots = 2
    env.provider.fail_showtimes = {"INGRESSO"}
    for _ in range(3):
        identifier = start_check(env, "failed", [301])
        assert env.app.state.jobs.snapshots[identifier].slices == {}


@pytest.mark.parametrize("path,expected", [
    ("/api/movies?country=BR", 3),
    ("/api/movie-groups?country=BR", 3),
    ("/api/provider-movies?country=BR", 3),
    ("/api/theaters?country=BR", 2),
    ("/api/provider-theaters?country=BR", 3),
    ("/api/format-mappings?country=BR", 3),
    ("/api/countries", 2),
    ("/api/connections?country=BR", 2),
])
def test_reference_and_catalog_lists_use_complete_server_pagination(env, path, expected):
    login(env)
    separator = "&" if "?" in path else "?"
    first = env.client.get(path + separator + "page=1&page_size=1")
    assert first.status_code == 200, first.text
    first = first.json()
    assert set(first) == {"items", "total", "page", "page_size"}
    assert first["total"] == expected and first["page_size"] == 1
    assert len(first["items"]) == 1 and first["page"] == 1
    second = env.client.get(path + separator + "page=2&page_size=1").json()
    assert len(second["items"]) == 1 and second["items"] != first["items"]
    exhausted = env.client.get(path + separator + f"page={expected + 1}&page_size=1").json()
    assert exhausted["items"] == [] and exhausted["total"] == expected
    assert env.provider.calls == []


def test_pagination_bounds_page_size_not_total_accessible_records(env):
    login(env)
    response = env.client.get("/api/provider-movies?page=-1&page_size=10000").json()
    assert response["page"] == 1 and response["page_size"] == 500 and response["total"] == 3
    assert len(response["items"]) == 3


def test_server_movie_search_supports_nested_provider_ids_and_exact_matches(env):
    login(env)
    for endpoint in ("movies", "movie-groups"):
        response = env.client.get(f"/api/{endpoint}?country=BR&q=ing-m&exact=true").json()
        assert response["total"] == 1 and response["items"][0]["id"] == 201
        assert response["items"][0]["provider_ids"][0]["external_id"] == "ing-m"
    response = env.client.get("/api/provider-movies?country=BR&q=ING-M&exact=true").json()
    assert response["total"] == 1 and response["items"][0]["id"] == 301
    assert env.client.get("/api/provider-movies?country=BR&q=Feature&exact=true").json()["total"] == 1
    assert env.client.get("/api/provider-movies?country=BR&q=Feature").json()["total"] == 2
    assert env.client.get("/api/provider-movies?country=BR&provider=CINEMARK").json()["total"] == 1
    assert env.client.get("/api/movies?country=BR&provider=CINEMARK").json()["items"][0]["provider_ids"][0]["provider_code"] == "CINEMARK"


def test_group_list_buckets_all_members_without_duplicates_and_retains_provider_filter(env):
    login(env)
    env.client.post("/api/movie-groups/move", json={"provider_movie_ids": [302], "target_group_id": 201})
    response = env.client.get("/api/movies?country=BR&page_size=1&q=Feature&exact=true").json()
    assert response["total"] == 1
    assert {member["id"] for member in response["items"][0]["provider_ids"]} == {301, 302}
    response = env.client.get("/api/movies?country=BR&provider=CINEMARK&q=Feature&exact=true").json()
    assert response["total"] == 1 and [member["id"] for member in response["items"][0]["provider_ids"]] == [302]


def test_reference_searches_name_ids_address_country_and_review_status_in_sql(env):
    login(env)
    response = env.client.get("/api/theaters?country=BR&q=Manual+address&exact=true").json()
    assert response["total"] == 1 and response["items"][0]["id"] == 101
    assert env.client.get("/api/theaters?country=BR&q=101&exact=true").json()["total"] == 1
    assert env.client.get("/api/theaters?country=BR&provider=CINEMARK").json()["total"] == 1
    assert env.client.get("/api/countries?q=BR&exact=true").json()["items"][0]["name"] == "Brazil"
    assert env.client.get("/api/countries?q=brazil&exact=true").json()["items"][0]["code"] == "BR"
    assert env.client.get("/api/countries?provider=CINEMARK").json()["total"] == 1
    assert env.client.get("/api/connections?q=brazil&exact=true").json()["total"] == 2
    response = env.client.get("/api/provider-theaters?country=BR&provider=CINEMARK&q=cin-t&exact=true&status=confirmed").json()
    assert response["total"] == 1 and response["items"][0]["theater_name"] == "Canonical Cinema"
    assert env.client.get("/api/provider-theaters?status=needs_review").json()["total"] == 0


def test_sql_search_treats_percent_and_underscore_as_literal_characters(env):
    login(env)
    assert env.client.get("/api/provider-movies?q=%25").json()["total"] == 0
    assert env.client.get("/api/provider-movies?q=_").json()["total"] == 0
    assert env.client.get("/api/theaters?q=%25").json()["total"] == 0


def test_movie_specific_format_list_is_enriched_and_paginated(env):
    login(env)
    response = env.client.post("/api/movie-format-overrides", json={
        "group_id": 201, "format_mapping_id": 501, "published_format": "IMAX Special", "active": True})
    assert response.status_code == 200, response.text
    response = env.client.get("/api/movie-format-overrides?country=BR&provider=INGRESSO&q=Feature&exact=true&page_size=1")
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["total"] == 1
    row = payload["items"][0]
    assert row["group_name"] == "Feature" and row["country_code"] == "BR"
    assert row["raw_format"] == "IMAX | Dublado" and row["provider_code"] == "INGRESSO"


def test_search_ranks_exact_and_prefix_before_earlier_alphabetic_substring(env):
    login(env)
    with env.app.state.sessions() as db:
        for identifier, name in ((104, "A PLF Cinema"), (105, "PLF Cinema"), (106, "PLF")):
            db.add(Theater(id=identifier, country_code="BR", name=name, timezone="America/Sao_Paulo"))
        db.commit()
    rows = env.client.get("/api/theaters?country=BR&q=PLF&page_size=2").json()
    assert rows["total"] == 3
    assert [row["name"] for row in rows["items"]] == ["PLF", "PLF Cinema"]


def test_ingresso_url_key_is_returned_and_searchable_without_changing_selected_identity(env):
    login(env)
    with env.app.state.sessions() as db:
        db.get(ProviderMovie, 301).url_key = "feature-url-key"
        db.commit()
    for endpoint in ("provider-movies", "movies", "movie-groups"):
        response = env.client.get(f"/api/{endpoint}?country=BR&q=FEATURE-URL-KEY&exact=true")
        assert response.status_code == 200, response.text
        result = response.json()
        assert result["total"] == 1
        item = result["items"][0] if endpoint == "provider-movies" else result["items"][0]["provider_ids"][0]
        assert item["url_key"] == "feature-url-key" and item["external_id"] == "ing-m"
        assert item["id"] == 301
    # Compatibility remains available for integrations still searching numeric/event IDs.
    assert env.client.get("/api/provider-movies?q=ing-m&exact=true").json()["total"] == 1
    identifier = start_check(env, provider_ids=[301])
    assert ("check", "INGRESSO", ["ing-m"]) in env.provider.calls
    assert env.client.get(f"/api/checks/{identifier}").status_code == 200


def test_ingresso_catalog_omission_keeps_existing_alias_and_assignment_end_to_end(env):
    login(env)
    with env.app.state.sessions() as db:
        db.get(ProviderMovie, 301).url_key = "feature-url-key"
        db.commit()
    refresh(env, [1])
    row = env.client.get("/api/provider-movies?q=feature-url-key&exact=true").json()["items"][0]
    assert row["id"] == 301 and row["group_id"] == 201 and row["manually_grouped"] is True
    assert row["external_id"] == "ing-m" and row["url_key"] == "feature-url-key"


def test_theaters_default_to_active_and_paginate_before_enriching_all_provider_ids(env):
    login(env)
    with env.app.state.sessions() as db:
        db.get(Theater, 102).active = False
        db.add(ProviderTheater(id=404, connection_id=1, external_id="review-alias", name="Review alias",
                              theater_id=101, review_status="needs_review"))
        db.add(ProviderTheater(id=405, connection_id=2, external_id="ignored-alias", name="Ignored alias",
                              theater_id=101, review_status="ignored"))
        db.commit()
    result = env.client.get("/api/theaters?country=BR&page_size=1").json()
    assert result["total"] == 1 and len(result["items"]) == 1
    aliases = result["items"][0]["provider_ids"]
    assert {row["external_id"] for row in aliases} == {"ing-t", "cin-t", "review-alias", "ignored-alias"}
    assert {row["review_status"] for row in aliases} == {"confirmed", "needs_review", "ignored"}
    assert env.client.get("/api/theaters?country=BR&page=2&page_size=1").json()["items"] == []
    all_rows = env.client.get("/api/theaters?country=BR&include_inactive=true&page_size=1").json()
    assert all_rows["total"] == 2
    second = env.client.get("/api/theaters?country=BR&include_inactive=true&page=2&page_size=1").json()
    assert second["items"][0]["id"] == 102 and second["items"][0]["active"] is False
    # Filtering to a provider controls canonical rows, not enrichment identities.
    filtered = env.client.get("/api/theaters?country=BR&provider=CINEMARK").json()
    assert filtered["total"] == 1
    assert {row["external_id"] for row in filtered["items"][0]["provider_ids"]} == {row["external_id"] for row in aliases}


def test_theater_identity_search_is_exact_compatible_and_never_duplicates_canonical_rows(env):
    login(env)
    for query in ("ing-t", "cin-t", "CINEMARK"):
        result = env.client.get(f"/api/theaters?country=BR&q={query}&exact=true").json()
        assert result["total"] == 1 and result["items"][0]["id"] == 101
    assert env.client.get("/api/theaters?country=BR&q=INGRESSO&exact=true").json()["total"] == 2
    with env.app.state.sessions() as db:
        db.add(ProviderTheater(id=404, connection_id=1, external_id="ing-t-extra", name="Extra", theater_id=101, review_status="confirmed"))
        db.commit()
    result = env.client.get("/api/theaters?country=BR&q=ing-t").json()
    assert result["total"] == 1 and len(result["items"]) == 1
    assert len(result["items"][0]["provider_ids"]) == 3


def test_theater_provider_identity_enrichment_query_count_does_not_grow_per_row(env):
    login(env)
    statements = []
    def record(_, __, statement, ___, ____, _____):
        if statement.strip().lower().startswith("select"):
            statements.append(statement)
    event.listen(env.engine, "before_cursor_execute", record)
    try:
        assert env.client.get("/api/theaters?country=BR&page_size=1").status_code == 200
        first = len(statements)
        statements.clear()
        assert env.client.get("/api/theaters?country=BR&page_size=2").status_code == 200
        second = len(statements)
    finally:
        event.remove(env.engine, "before_cursor_execute", record)
    # Authentication, canonical count, canonical page, one association query.
    assert first == second == 4


def test_optional_is_catalog_window_is_validated_exposed_and_nullable(env):
    login(env)
    response = env.client.post("/api/connections", json={"country_code": "AU", "provider_code": "IS",
        "provider_country_code": "AU", "timezone": "Australia/Sydney", "catalog_window_days": 30})
    assert response.status_code == 200, response.text
    identifier = response.json()["id"]
    assert response.json()["catalog_window_days"] == 30
    assert next(row for row in env.client.get("/api/bootstrap").json()["connections"] if row["id"] == identifier)["catalog_window_days"] == 30
    for value, expected in ((None, None), (0, None), (1, 1), (3650, 3650)):
        response = env.client.patch(f"/api/connections/{identifier}", json={"catalog_window_days": value})
        assert response.status_code == 200 and response.json()["catalog_window_days"] == expected
    for value in (-1, 3651, True, 1.2, "30"):
        assert env.client.patch(f"/api/connections/{identifier}", json={"catalog_window_days": value}).status_code == 400
    assert env.client.patch("/api/connections/1", json={"catalog_window_days": 30}).status_code == 400
    assert env.provider.calls == []


def test_explicit_archived_provider_movie_can_still_load_fresh_sessions(env):
    login(env)
    with env.app.state.sessions() as db:
        db.get(ProviderMovie, 301).active = False
        db.get(MovieGroup, 201).active = False
        db.commit()
    assert env.client.get("/api/provider-movies?q=ing-m&exact=true").json()["total"] == 0
    assert env.client.get("/api/provider-movies?q=ing-m&exact=true&include_inactive=true").json()["total"] == 1
    identifier = start_check(env, provider_ids=[301])
    assert ("check", "INGRESSO", ["ing-m"]) in env.provider.calls
    assert env.client.get(f"/api/checks/{identifier}").json()["total"] == 3


def test_malformed_is_catalog_preserves_existing_active_ids_and_manual_groups(env):
    from app.providers import ProviderService
    login(env)
    with env.app.state.sessions() as db:
        db.add(Connection(id=3, country_code="AU", provider_code="IS", provider_country_code="AU", timezone="Australia/Sydney"))
        db.add(MovieGroup(id=204, country_code="AU", name="Existing IS movie"))
        db.flush()
        db.add(ProviderMovie(id=304, connection_id=3, external_id="existing-id", name="Existing IS movie",
                             group_id=204, active=True, manually_grouped=True))
        db.commit()
    service = ProviderService(SimpleNamespace(is_api_key="fixture-key", request_timeout=20))
    service._fetch_json = lambda *_args, **_kwargs: {"movies": [{"title": "Missing native ID"}], "meta_info": {"total_count": 1}}
    env.provider.catalog = service.catalog
    refresh(env, [3], "failed")
    with env.app.state.sessions() as db:
        movie = db.get(ProviderMovie, 304)
        assert movie.active is True and movie.group_id == 204 and movie.manually_grouped is True
        assert db.get(MovieGroup, 204).active is True


def test_is_untitled_catalog_keeps_known_metadata_and_discovers_new_native_ids(env):
    from datetime import date
    from app.providers import ProviderService
    login(env)
    with env.app.state.sessions() as db:
        db.add(Connection(id=3, country_code="AU", provider_code="IS", provider_country_code="AU", timezone="Australia/Sydney"))
        db.add(MovieGroup(id=204, country_code="AU", name="Manual display group"))
        db.flush()
        db.add(ProviderMovie(id=304, connection_id=3, external_id="existing-id", name="Known display title",
                             original_name="Known original title", release_date=date(1991, 1, 2),
                             group_id=204, active=False, manually_grouped=True))
        db.commit()
    service = ProviderService(SimpleNamespace(is_api_key="fixture-key", request_timeout=20))
    service._fetch_json = lambda *_args, **_kwargs: {
        "movies": [{"id": "existing-id", "title": "", "original_title": None, "release_dates": {"AU": "2099-01-01"}},
                   {"id": "new-id", "title": None}], "meta_info": {"total_count": 2}}
    env.provider.catalog = service.catalog
    refresh(env, [3], "done")
    with env.app.state.sessions() as db:
        known = db.get(ProviderMovie, 304)
        assert known.name == "Known display title" and known.original_name == "Known original title"
        assert known.release_date == date(1991, 1, 2) and known.active is True
        assert known.group_id == 204 and known.manually_grouped is True
        assert db.get(MovieGroup, 204).name == "Manual display group"
        new = db.scalar(select(ProviderMovie).where(ProviderMovie.connection_id == 3, ProviderMovie.external_id == "new-id"))
        assert new.name == "IS movie new-id" and new.active is True
        new_id, new_group = new.id, new.group_id
        assert db.get(Connection, 3).movie_count == 2
    visible = env.client.get("/api/provider-movies?country=AU&q=IS%20movie%20new-id&exact=true").json()
    assert visible["total"] == 1 and visible["items"][0]["id"] == new_id
    service._fetch_json = lambda *_args, **_kwargs: {
        "movies": [{"id": "existing-id", "title": "", "original_title": None},
                   {"id": "new-id", "title": "Newly supplied title", "original_title": "New original"}],
        "meta_info": {"total_count": 2}}
    refresh(env, [3], "done")
    with env.app.state.sessions() as db:
        new = db.get(ProviderMovie, new_id)
        assert new.name == "Newly supplied title" and new.original_name == "New original"
        assert new.group_id == new_group
        assert db.get(ProviderMovie, 304).name == "Known display title"


def test_incomplete_ingresso_pagination_preserves_existing_catalog(env):
    from app.providers import ProviderService
    login(env)
    service = ProviderService(SimpleNamespace(ingresso_partnerships=["warner"], request_timeout=20))
    service._fetch_json = lambda *_args, **_kwargs: {"items": [{"id": "ing-m", "title": "Film"}], "count": 1}
    env.provider.catalog = service.catalog
    refresh(env, [1], "failed")  # A non-advancing repeated page, never full archive fallback.
    with env.app.state.sessions() as db:
        assert db.get(ProviderMovie, 301).active is True and db.get(ProviderMovie, 303).active is True
        assert db.get(ProviderMovie, 301).group_id == 201 and db.get(ProviderMovie, 301).manually_grouped is True


def test_archived_movie_picker_excludes_disabled_connections_without_losing_admin_reference_access(env):
    login(env)
    with env.app.state.sessions() as db:
        db.get(Connection, 1).enabled = False
        db.get(ProviderMovie, 301).active = False
        db.commit()
    for suffix in ("", "&include_inactive=true"):
        response = env.client.get("/api/provider-movies?q=ing-m&exact=true" + suffix)
        assert response.status_code == 200 and response.json()["total"] == 0
    assert env.client.get("/api/provider-movies?include_inactive=true").json()["total"] == 1
    references = env.client.get("/api/movie-groups?country=BR&include_inactive=true&q=ing-m&exact=true").json()
    assert references["total"] == 1 and references["items"][0]["provider_ids"][0]["id"] == 301
    response = env.client.post("/api/checks", json={"country_code": "BR", "provider_movie_ids": [301]})
    assert response.status_code == 400
    assert env.provider.calls == []
