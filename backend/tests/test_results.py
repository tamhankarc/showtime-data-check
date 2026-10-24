"""Regression checks for safe exports, source VIP distinctions, and worker payload release."""
import gc
from io import BytesIO
import json
import tempfile
import threading
import time
import weakref

from openpyxl import load_workbook
from fastapi.testclient import TestClient
import pytest
from sqlalchemy.orm import sessionmaker

from app.config import Settings
from app.db import make_engine
from app.jobs import JobManager
from app.models import Base, Connection, Country, FormatMapping, MovieGroup, ProviderMovie, ProviderTheater, Theater
from app.providers import ProviderResult
from app.results import consolidate, normalize_rows, xlsx_export
from app.security import password_hash


@pytest.fixture
def reference_database(tmp_path):
    engine = make_engine("sqlite:///" + (tmp_path / "results.sqlite").as_posix())
    Base.metadata.create_all(engine)
    sessions = sessionmaker(bind=engine, expire_on_commit=False)
    with sessions() as db:
        db.add_all([Country(code="BR", name="Brazil", timezone="America/Sao_Paulo"),
                    Country(code="AU", name="Australia", timezone="Australia/Sydney"),
                    Country(code="DE", name="Germany", timezone="Europe/Berlin")])
        db.flush()
        db.add_all([Connection(id=1, country_code="BR", provider_code="INGRESSO", provider_country_code="BR", timezone="America/Sao_Paulo"),
                    Connection(id=2, country_code="BR", provider_code="CINEMARK", provider_country_code="BR", timezone="America/Sao_Paulo"),
                    Connection(id=3, country_code="AU", provider_code="IS", provider_country_code="AU", timezone="Australia/Sydney"),
                    MovieGroup(id=1, country_code="BR", name="Fixture feature"),
                    Theater(id=1, country_code="BR", name="Confirmed theater", timezone="America/Sao_Paulo")])
        db.flush()
        db.add_all([ProviderMovie(id=1, connection_id=2, external_id="movie", name="Fixture feature", group_id=1),
                    ProviderTheater(id=1, connection_id=2, external_id="cinema", name="Provider theater", theater_id=1, review_status="confirmed"),
                    FormatMapping(id=1, country_code="BR", provider_code="CINEMARK", raw_format="IMAX", published_format="IMAX")])
        db.commit()
    yield engine, sessions
    engine.dispose()


@pytest.fixture
def authenticated_app(reference_database):
    from app.main import create_app
    engine, _ = reference_database
    password = "independent-fixture-password"
    settings = Settings(proxy_secret="fixture-private-proxy", admin_password_hash=password_hash(password))
    app = create_app(settings, engine=engine)
    with TestClient(app) as client:
        client.headers.update({"x-proxy-secret": settings.proxy_secret, "origin": settings.app_origin})
        assert client.post("/api/login", json={"username": settings.admin_username, "password": password}).status_code == 200
        yield client


@pytest.mark.parametrize("identifier,payload", [(1, {"provider_code": "VELOX"}), (3, {"country_code": "DE"})])
def test_existing_connection_provider_and_country_identities_cannot_be_changed(
    reference_database, authenticated_app, identifier, payload
):
    response = authenticated_app.patch(f"/api/connections/{identifier}", json=payload)
    assert response.status_code == 400
    _, sessions = reference_database
    with sessions() as db:
        connection = db.get(Connection, identifier)
        assert (connection.country_code, connection.provider_code) == (("BR", "INGRESSO") if identifier == 1 else ("AU", "IS"))


def test_canonical_theater_cannot_move_out_of_its_confirmed_provider_country(reference_database, authenticated_app):
    response = authenticated_app.patch("/api/theaters/1", json={"country_code": "AU"})
    assert response.status_code == 400
    _, sessions = reference_database
    with sessions() as db:
        assert db.get(Theater, 1).country_code == "BR"
        assert db.get(ProviderTheater, 1).theater_id == 1


@pytest.mark.parametrize("view", ["showtimes", "theaters"])
def test_export_strips_xml_controls_before_formula_protection_and_preserves_valid_text(view):
    row = {"movie_name": '\x00=HYPERLINK("https://example.invalid")\x01',
           "theater_name": "Cinéma\x08\x0b\x0c\x0e\x1f 東京",
           "address": "\tStreet\nSecond line\rEnd", "raw_format": "IMAX\x02"}
    content = xlsx_export([row], {"note": "\x01+SUM(1,1)"}, view)
    book = load_workbook(BytesIO(content), data_only=False)
    sheet = book["Theaters" if view == "theaters" else "Showtimes"]
    assert sheet["A2"].value == '\'=HYPERLINK("https://example.invalid")'
    assert sheet["A2"].data_type != "f"
    assert sheet["C2"].value == "Cinéma 東京"
    assert sheet["E2"].value == "\tStreet\nSecond line\nEnd"
    assert book["Check details"]["B2"].value == "'+SUM(1,1)"
    assert book["Check details"]["B2"].data_type != "f"


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_export_provider_values_never_become_workbook_formulas(prefix):
    content = xlsx_export([{"movie_name": prefix + "provider text"}], {}, "showtimes")
    cell = load_workbook(BytesIO(content), data_only=False)["Showtimes"]["A2"]
    assert cell.value == "'" + prefix + "provider text"
    assert cell.data_type != "f"


@pytest.mark.parametrize("view", ["showtimes", "theaters"])
def test_export_empty_strings_remain_blank_instead_of_becoming_apostrophes(view):
    content = xlsx_export([{"movie_name": "", "theater_name": ""}], {"empty": ""}, view)
    book = load_workbook(BytesIO(content), data_only=False)
    sheet = book["Theaters" if view == "theaters" else "Showtimes"]
    assert sheet["A2"].value is None and sheet["C2"].value is None
    assert book["Check details"]["B2"].value is None


@pytest.mark.parametrize("view", ["showtimes", "theaters"])
def test_xlsx_exports_never_create_temporary_files_and_keep_workbook_contract(monkeypatch, view):
    def reject_temporary_file(*args, **kwargs):
        raise AssertionError("Export bodies must never be written to a temporary file")

    monkeypatch.setattr(tempfile, "mkstemp", reject_temporary_file)
    row = {"movie_name": "In-memory feature", "theater_name": "https://cinema.example/movie",
           "country_code": "BR", "providers": ["CINEMARK"], "showtimes": 2}
    metadata = {"status": "partial", "selection": {"country_code": "BR", "provider_codes": ["CINEMARK"]},
                "provider_outcomes": [{"provider_code": "CINEMARK", "fetched_at": "2026-10-24T12:00:00+00:00"}]}
    content = xlsx_export([row], metadata, view)
    book = load_workbook(BytesIO(content), data_only=False)
    name = "Theaters" if view == "theaters" else "Showtimes"
    assert book.sheetnames == [name, "Check details"]
    sheet = book[name]
    assert sheet["A2"].value == "In-memory feature"
    assert sheet["C2"].value == "https://cinema.example/movie" and sheet["C2"].hyperlink is None
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == ("A1:L2" if view == "theaters" else "A1:Q2")
    assert sheet["A1"].font.bold and sheet["A1"].fill.fgColor.rgb.endswith("173B63")
    details = dict(book["Check details"].iter_rows(min_row=2, values_only=True))
    assert details["status"] == "partial"
    assert json.loads(details["selection"]) == metadata["selection"]
    assert json.loads(details["provider_outcomes"]) == metadata["provider_outcomes"]
    assert book["Check details"]["B2"].alignment.wrap_text


def test_provider_vip_flag_keeps_same_time_canonical_screenings_separate(reference_database):
    _, sessions = reference_database
    common = {"title": "movie", "theater": "cinema", "date": "2099-10-24", "time": "14:00",
              "format": "IMAX", "audio": "EN", "subtitle": "", "links": []}
    with sessions() as db:
        connection = db.get(Connection, 2)
        rows = normalize_rows(db, connection, [
            {**common, "id": "standard", "screen": "Screen 1", "vip": False},
            {**common, "id": "premium", "screen": "VIP Screen", "vip": True},
        ])
    assert {row["vip"] for row in rows} == {False, True}
    assert all(row["published_format"] == "IMAX" and row["raw_format"] == "IMAX" for row in rows)
    assert len(consolidate(rows)) == 2


@pytest.mark.parametrize("failure_stage", [None, "normalization", "provider"])
def test_worker_releases_raw_payloads_and_expired_rows_while_another_provider_waits(
    reference_database, monkeypatch, failure_stage
):
    _, sessions = reference_database
    release = threading.Event()
    refs = {}

    class Body(dict):
        """Weak references detect retained showtime bodies without inspecting frame internals."""

    class Provider:
        def showtimes(self, connection, ids, cancel):
            if connection["provider_code"] == "CINEMARK":
                # Intentionally wait independently of cancellation to expose lingering worker aliases.
                if not release.wait(10):
                    raise RuntimeError("Fixture did not release the second provider")
                return ProviderResult()
            body = Body(synthetic_showtime="raw fixture body")
            refs["raw"] = weakref.ref(body)
            if failure_stage == "provider":
                raise ValueError("Synthetic provider parsing failure")
            return ProviderResult(showtimes=[body])

    def normalize(db, connection, source):
        if connection.provider_code != "INGRESSO":
            return []
        if failure_stage == "normalization":
            raise ValueError("Synthetic normalization failure")
        body = Body(synthetic_showtime="normalized fixture body")
        refs["normalized"] = weakref.ref(body)
        return [body]

    monkeypatch.setattr("app.jobs.normalize_rows", normalize)
    manager = JobManager(sessions, Settings(snapshot_seconds=900), Provider())
    try:
        identifier = manager.submit("check", "fixture-viewer", {
            "country_code": "BR", "connection_ids": [1, 2],
            "external_ids": {"INGRESSO": ["movie"], "CINEMARK": ["movie"]},
        })
        until = time.monotonic() + 5
        processed = False
        while time.monotonic() < until:
            with manager.lock:
                snapshot = manager.snapshots[identifier]
                processed = "INGRESSO" in snapshot.outcomes
            if processed:
                break
            time.sleep(0.01)
        assert processed, "First provider should finish while the second remains blocked"
        gc.collect()
        assert refs["raw"]() is None
        if failure_stage is None:
            assert refs["normalized"]() is not None
            with manager.lock:
                manager.snapshots[identifier].slices["INGRESSO"].deadline = time.monotonic() - 1
            manager.purge()
            gc.collect()
            assert refs["normalized"]() is None
    finally:
        release.set()
        manager.close()
        manager.executor.shutdown(wait=True, cancel_futures=True)
